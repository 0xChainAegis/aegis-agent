"""导出阶段总结为 Word。依赖：py -m pip install python-docx==1.2.0
在仓库根目录执行：py scripts/export-work-summary.py
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = root / 'docs/LLM阶段工作总结.md'
output = root / 'docs/AegisAgent_LLM阶段工作总结.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2.2)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.2
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].font.color.rgb = RGBColor.from_string('24476A')

def clean(text):
    return text.replace('**', '').replace('`', '').replace('  ', ' ').strip()

lines = source.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [clean(c) for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', c.replace(' ', '')) for c in cells):
                rows.append(cells)
            i += 1
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = 'Table Grid'
        for n, cells in enumerate(rows):
            row = table.add_row()
            for cell, text in zip(row.cells, cells):
                cell.text = text
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        run.bold = n == 0
                if n == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'E8EFF6')
                    cell._tc.get_or_add_tcPr().append(shade)
        doc.add_paragraph()
        continue
    if line.startswith('# '):
        doc.add_heading(clean(line[2:]), 0)
    elif line.startswith('### '):
        doc.add_heading(clean(line[4:]), 2)
    elif line.startswith('## '):
        doc.add_heading(clean(line[3:]), 1)
    elif line.startswith('- '):
        doc.add_paragraph(clean(line[2:]), style='List Bullet')
    elif line:
        doc.add_paragraph(clean(line))
    i += 1
sec.footer.paragraphs[0].text = 'AegisAgent · LLM 模块阶段工作总结'
doc.core_properties.author = '满天星斗烟火繁'
doc.core_properties.title = 'AegisAgent LLM 模块阶段工作总结'
doc.save(output)
# 重新读取并验证核心内容及表格，确保交付物可解析且数据未丢失。
check = Document(output)
text = '\n'.join(p.text for p in check.paragraphs) + '\n' + '\n'.join(c.text for t in check.tables for r in t.rows for c in r.cells)
for token in ['40', '120', '85.0%', '67.5%', '30.0%', '产品级最终正确率', 'PR #7']:
    assert token in text, token
assert len(check.tables) == 6, len(check.tables)
print('Verified Word document:', output, 'bytes=', output.stat().st_size)
